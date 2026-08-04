from pathlib import Path

import pypdfium2 as pdfium
import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.services.confirmations.extract import (
    MIN_TEXT_CHARS_PER_PAGE, extract_document,
)


def _write_text_pdf(path: Path) -> None:
    # Create a minimal valid PDF with text content
    # This PDF has a single page with text drawn via PDF text operators
    text = "Trade Confirmation: BUY 100 EuropeanVanillaOption on AAPL strike 150 maturity 1.0 premium 12.5 USD"

    # Build PDF content stream (the text drawing commands)
    content_stream = b"BT /F1 12 Tf 50 700 Td (" + text.encode('latin-1') + b") Tj ET\n"

    # Build the complete PDF with proper structure
    pdf_parts = [
        b"%PDF-1.4\n",
        # Object 1: Catalog
        b"1 0 obj\n<</Type/Catalog/Pages 2 0 R>>\nendobj\n",
        # Object 2: Pages
        b"2 0 obj\n<</Type/Pages/Kids[3 0 R]/Count 1>>\nendobj\n",
        # Object 3: Page
        b"3 0 obj\n<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>\nendobj\n",
        # Object 4: Content stream
        b"4 0 obj\n<</Length " + str(len(content_stream)).encode() + b">>\nstream\n" + content_stream + b"endstream\nendobj\n",
        # Object 5: Font
        b"5 0 obj\n<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>\nendobj\n",
    ]

    # Calculate byte offsets for xref table
    offsets = []
    current_offset = len(pdf_parts[0])  # Start after %PDF-1.4\n
    offsets.append(0)  # Placeholder for object 0

    for i, part in enumerate(pdf_parts[1:], 1):
        offsets.append(current_offset)
        current_offset += len(part)

    # Add xref and trailer
    xref_offset = current_offset
    xref_lines = [b"xref\n0 " + str(len(offsets)).encode() + b"\n"]
    xref_lines.append(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        xref_lines.append(f"{offset:010d} 00000 n \n".encode())

    trailer = b"trailer<</Size " + str(len(offsets)).encode() + b"/Root 1 0 R>>\n"
    startxref = b"startxref\n" + str(xref_offset).encode() + b"\n%%EOF\n"

    # Write complete PDF
    pdf_content = b"".join(pdf_parts) + b"".join(xref_lines) + trailer + startxref
    path.write_bytes(pdf_content)


def _write_blank_pdf(path: Path) -> None:
    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    with path.open("wb") as fh:
        w.write(fh)


def _write_docx(path: Path, text: str) -> None:
    doc = DocxDocument()
    doc.add_paragraph(text)
    doc.save(str(path))


def test_text_pdf_extracts_text_mode(tmp_path):
    p = tmp_path / "text.pdf"
    _write_text_pdf(p)
    content = extract_document(p)
    assert content.extract_mode == "text"
    assert content.page_count == 1
    assert "EuropeanVanillaOption" in content.pages[0].text
    assert content.pages[0].image_png is None


def test_blank_pdf_page_renders_image(tmp_path):
    p = tmp_path / "scan.pdf"
    _write_blank_pdf(p)
    content = extract_document(p)
    assert content.extract_mode == "vision"
    assert content.pages[0].image_png is not None
    assert content.pages[0].image_png[:8] == b"\x89PNG\r\n\x1a\n"


def test_docx_extracts_paragraph_and_table_text(tmp_path):
    p = tmp_path / "conf.docx"
    doc = DocxDocument()
    doc.add_paragraph("Confirmation of SnowballOption trade")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Underlying"
    table.rows[0].cells[1].text = "700050.SH"
    doc.save(str(p))
    content = extract_document(p)
    assert content.extract_mode == "text"
    assert "SnowballOption" in content.pages[0].text
    assert "700050.SH" in content.pages[0].text


def test_unsupported_extension_raises(tmp_path):
    p = tmp_path / "conf.txt"
    p.write_text("hello")
    with pytest.raises(ValueError, match="Unsupported"):
        extract_document(p)


def test_docx_mixed_text_and_images(tmp_path):
    """Test DOCX with near-empty text + embedded image uses 1-based indexes."""
    from PIL import Image

    p = tmp_path / "mixed.docx"
    doc = DocxDocument()

    # Add minimal text (< MIN_TEXT_CHARS_PER_PAGE)
    doc.add_paragraph("Short")

    # Create and embed a tiny PNG image
    img = Image.new("RGB", (10, 10), color="red")
    img_path = tmp_path / "tiny.png"
    img.save(img_path)
    doc.add_picture(str(img_path))

    doc.save(str(p))

    # Extract and verify
    content = extract_document(p)

    # Should be mixed or vision mode (text too short to be sole page)
    assert content.extract_mode in ("mixed", "vision")
    assert content.page_count >= 1

    # All page indexes must be 1-based and consecutive
    indexes = [page.index for page in content.pages]
    assert indexes == list(range(1, len(indexes) + 1)), f"Indexes not 1-based consecutive: {indexes}"

    # At least one page must have PNG image data
    image_pages = [p for p in content.pages if p.image_png is not None]
    assert len(image_pages) > 0, "No image pages found"

    # Image bytes must be valid PNG
    for img_page in image_pages:
        assert img_page.image_png[:8] == b"\x89PNG\r\n\x1a\n", "Invalid PNG magic bytes"
