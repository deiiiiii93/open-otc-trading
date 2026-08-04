"""PDF/DOCX text + image extraction for trade confirmations.

Text-first: a PDF page yielding fewer than MIN_TEXT_CHARS_PER_PAGE characters is
treated as a scan and rendered to PNG (pypdfium2) for the vision extraction
path. DOCX is page-less: all paragraphs + table cells become one text "page";
if that text is under the threshold and the archive embeds images, those images
are surfaced for vision instead.
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

MIN_TEXT_CHARS_PER_PAGE = 40
_RENDER_SCALE = 2.0  # ~144 dpi; keeps scans legible without huge payloads


@dataclass(frozen=True)
class PageContent:
    index: int
    text: str
    image_png: bytes | None = None


@dataclass(frozen=True)
class DocumentContent:
    pages: list[PageContent]
    page_count: int
    extract_mode: str


def extract_document(path: Path) -> DocumentContent:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    raise ValueError(f"Unsupported confirmation file type: {suffix!r} (want .pdf/.docx)")


def _mode(pages: list[PageContent]) -> str:
    has_text = any(p.image_png is None for p in pages)
    has_image = any(p.image_png is not None for p in pages)
    if has_text and has_image:
        return "mixed"
    return "vision" if has_image else "text"


def _extract_pdf(path: Path) -> DocumentContent:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    texts = [(page.extract_text() or "").strip() for page in reader.pages]
    pages: list[PageContent] = []
    image_indexes = [
        i for i, text in enumerate(texts) if len(text) < MIN_TEXT_CHARS_PER_PAGE
    ]
    rendered: dict[int, bytes] = {}
    if image_indexes:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(path))
        try:
            for i in image_indexes:
                bitmap = pdf[i].render(scale=_RENDER_SCALE)
                pil = bitmap.to_pil()
                buf = io.BytesIO()
                pil.save(buf, format="PNG")
                rendered[i] = buf.getvalue()
        finally:
            pdf.close()
    for i, text in enumerate(texts):
        pages.append(PageContent(index=i + 1, text=text, image_png=rendered.get(i)))
    return DocumentContent(pages=pages, page_count=len(pages), extract_mode=_mode(pages))


def _extract_docx(path: Path) -> DocumentContent:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    chunks = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    text = "\n".join(chunks).strip()
    if len(text) >= MIN_TEXT_CHARS_PER_PAGE:
        return DocumentContent(
            pages=[PageContent(index=1, text=text)], page_count=1, extract_mode="text"
        )
    # Near-empty text: surface embedded images (a scan pasted into Word).
    images: list[bytes] = []
    with zipfile.ZipFile(str(path)) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/") and name.lower().endswith(
                (".png", ".jpg", ".jpeg")
            ):
                images.append(zf.read(name))
    if not images:
        return DocumentContent(
            pages=[PageContent(index=1, text=text)], page_count=1, extract_mode="text"
        )
    pages = [
        PageContent(index=i + 1, text="", image_png=_ensure_png(data))
        for i, data in enumerate(images)
    ]
    if text:
        pages.insert(0, PageContent(index=0, text=text))
    return DocumentContent(pages=pages, page_count=len(pages), extract_mode=_mode(pages))


def _ensure_png(data: bytes) -> bytes:
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return data
    from PIL import Image

    buf = io.BytesIO()
    Image.open(io.BytesIO(data)).convert("RGB").save(buf, format="PNG")
    return buf.getvalue()
